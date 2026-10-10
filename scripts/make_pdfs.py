r"""Manual smoke check: extract text from a real .docx, build a quiz, render PDFs.

Not part of the test suite. Run from the project root with:
    .venv\Scripts\python scripts\make_pdfs.py
"""

import os
import sys
import tempfile
from pathlib import Path

import django
from django.contrib.auth import get_user_model

# Allow running this file directly by putting the project root on sys.path.
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "quizcraft.settings")
os.chdir(ROOT)
django.setup()

from quizzes.models import (  # noqa: E402
    Difficulty,
    Question,
    QuestionType,
    Quiz,
    SourceDocument,
)
from quizzes.services import export  # noqa: E402
from quizzes.services.extraction import extract_text  # noqa: E402

USER = "smoke"
TITLE = "Mitochondria Mastery"

BODY = (
    "Mitochondria generate most of the cell's ATP through oxidative phosphorylation. "
    "The inner membrane folds into structures called cristae, which expand the "
    "surface area available for the electron transport chain to operate on. "
    "Protons accumulate in the intermembrane space and flow back through ATP "
    "synthase, and that flow is what actually drives the production of ATP. "
    "Mitochondria also carry their own circular DNA and ribosomes, which is why "
    "they are described as semi-autonomous organelles."
)

SPECS = [
    (
        QuestionType.MCQ_SINGLE,
        "Which organelle produces the majority of cellular ATP?",
        ["Nucleus", "Mitochondrion", "Ribosome", "Golgi apparatus"],
        [1],
        [],
        "Mitochondria generate ATP via oxidative phosphorylation.",
    ),
    (
        QuestionType.MCQ_MULTI,
        "Select all structures found inside a mitochondrion.",
        ["Cristae", "Nucleolus", "Thylakoid", "Matrix"],
        [0, 3],
        [],
        "Both the cristae and the matrix are internal to the mitochondrion.",
    ),
    (
        QuestionType.TRUE_FALSE,
        "Mitochondria contain their own DNA.",
        ["True", "False"],
        [0],
        ["true"],
        "Mitochondria are semi-autonomous and carry their own genome.",
    ),
    (
        QuestionType.FILL_BLANK,
        "The inner membrane of a mitochondrion folds into _____ to expand surface area.",
        [],
        [0],
        ["cristae", "cristae (inner folds)"],
        "Cristae increase the area available for the electron transport chain.",
    ),
    (
        QuestionType.SHORT_ANSWER,
        "Name the process that drives ATP synthesis in mitochondria.",
        [],
        [0],
        ["chemiosmosis", "chemiosmotic coupling"],
        "Protons flowing down their gradient through ATP synthase drives it.",
    ),
]


def make_docx_upload():
    """A real .docx, so the extraction path runs for real."""
    import io

    import docx
    from django.core.files.uploadedfile import SimpleUploadedFile

    document = docx.Document()
    document.add_heading("Cell Biology Notes", level=1)
    document.add_paragraph(BODY)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Structure"
    table.cell(0, 1).text = "Function"
    table.cell(1, 0).text = "Cristae"
    table.cell(1, 1).text = "Expand surface area for the ETC"

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return SimpleUploadedFile(
        "cell-biology.docx",
        buffer.read(),
        content_type=(
            "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.document"
        ),
    )


def main():
    user, _ = get_user_model().objects.get_or_create(username=USER)

    upload = make_docx_upload()
    # extract_text() works on a real path, so write the upload out first.
    with tempfile.TemporaryDirectory() as tmp:
        staged = Path(tmp) / upload.name
        with open(staged, "wb") as handle:
            handle.write(upload.read())
        text, page_count = extract_text(staged, upload.name)
    print(f"extracted {len(text)} chars from .docx (pages={page_count})")
    assert "cristae" in text.lower(), "extraction lost the body text"

    document = SourceDocument.objects.create(
        user=user,
        file=upload,
        original_filename="cell-biology.docx",
        file_size=upload.size,
        char_count=len(text),
        extracted_text=text,
    )

    Quiz.objects.filter(user=user, title=TITLE).delete()
    quiz = Quiz.objects.create(
        user=user,
        document=document,
        title=TITLE,
        question_count=len(SPECS),
        difficulty=Difficulty.MEDIUM,
        status=Quiz.Status.READY,
        question_types=list(QuestionType.values),
    )
    for order, (qtype, prompt, choices, key, answers, explanation) in enumerate(
        SPECS, 1
    ):
        question = Question.objects.create(
            quiz=quiz,
            order=order,
            question_type=qtype,
            prompt=prompt,
            choices=choices,
            correct_choices=key,
            acceptable_answers=answers,
            explanation=explanation,
            points=2,
        )
        assert question.type_label, "type_label rendered empty"
        print(f"  {order}. [{question.type_label}] {prompt[:52]}")

    for reveal, label in ((False, "student"), (True, "answer-key")):
        data = export.build_quiz_pdf(quiz, reveal_answers=reveal)
        name = export.safe_filename(f"{quiz.title}-{label}")
        with open(name, "wb") as handle:
            handle.write(data)
        print(f"wrote {name} ({len(data)} bytes)")

    print("quiz id:", quiz.pk)


if __name__ == "__main__":
    main()
